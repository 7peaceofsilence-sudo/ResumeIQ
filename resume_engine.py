import json
from pathlib import Path
from typing import Any

from groq import Groq
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document


class ResumeEngineError(Exception):
    """Readable error raised while extracting or analyzing a resume."""


class JobD(BaseModel):
    role: str = ""
    required_skills: list[str] = Field(default_factory=list)
    preferred_skills: list[str] = Field(default_factory=list)
    minimum_experience: float | None = None
    education_requirements: list[str] = Field(default_factory=list)
    responsibilities: list[str] = Field(default_factory=list)


class MatchResult(BaseModel):
    score: float = Field(ge=0, le=100)
    details: dict[str, Any] = Field(default_factory=dict)


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = Field(default_factory=list)


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience_years: float | None = None
    skills: list[str] = Field(default_factory=list)
    experiences: list[Experience] = Field(default_factory=list)
    education: list[str] = Field(default_factory=list)
    projects: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)


def extract_text_from_file(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(file_path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        document = Document(str(file_path))
        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    raise ResumeEngineError("Unsupported file type. Please use PDF or DOCX.")


class GroqResumeEngine:
    def __init__(self, api_key: str, model: str = "openai/gpt-oss-120b"):
        if not api_key:
            raise ResumeEngineError("GROQ_API_KEY is missing.")
        self.client = Groq(api_key=api_key)
        self.model = model

    def _json_completion(self, messages: list[dict[str, str]]) -> dict[str, Any]:
        try:
            response = self.client.chat.completions.create(
                model=self.model,
                messages=messages,
                response_format={"type": "json_object"},
            )
            content = response.choices[0].message.content
            if not content:
                raise ResumeEngineError("The AI model returned an empty response.")
            return json.loads(content)
        except json.JSONDecodeError as exc:
            raise ResumeEngineError("The AI response was not valid JSON. Please retry.") from exc
        except ResumeEngineError:
            raise
        except Exception as exc:
            # Keep error useful while avoiding a traceback in the web page.
            raise ResumeEngineError(f"Groq request failed: {exc}") from exc

    def parse_job_description(self, job_description: str) -> JobD:
        schema = JobD.model_json_schema()
        system_prompt = f"""
You are an expert HR assistant. Extract facts from the job description and return
ONLY a JSON object matching this schema:
{json.dumps(schema)}
Do not return the schema itself. Do not invent information.
If minimum experience is not mentioned, return null.
If list information is missing, return an empty list.
"""
        data = self._json_completion([
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": "Analyze this job description:\n\n" + job_description},
        ])
        return JobD(**data)

    def parse_resume(self, resume_text: str) -> Resume:
        schema = Resume.model_json_schema()
        system_prompt = f"""
You are an expert resume parser. Extract information by meaning, not just exact
section headings. Work history may appear under Experience, Professional Experience,
Work History, Employment or Internships. Extract skills across the resume, including
projects and internships. Include internships in experiences.
Return ONLY valid JSON matching this schema:
{json.dumps(schema)}
Do not invent information. Use null for unavailable scalar values and [] for unavailable lists.
"""
        data = self._json_completion([
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": "Parse the following resume:\n\n" + resume_text},
        ])
        return Resume(**data)

    def final_score(self, job: JobD, resume: Resume) -> MatchResult:
        schema = MatchResult.model_json_schema()
        prompt = f"""
You are an HR assistant comparing a resume with a job description.

JOB DESCRIPTION STRUCTURE:
{job.model_dump_json(indent=2)}

CANDIDATE RESUME:
{resume.model_dump_json(indent=2)}

Return only a JSON object matching this schema:
{json.dumps(schema)}

In details, include:
- candidate_name
- matching_skills (list)
- missing_important_skills (list)
- experience_requirement_met (boolean, string, or null when not applicable)
- overall_match_percentage (number from 0 to 100)
- final_verdict (short string)
- improvement_suggestions (list of concise, practical suggestions)

Keep the score evidence-based. Do not infer qualifications that are not supported by the resume.
The top-level score must be from 0 to 100.
"""
        data = self._json_completion([
            {"role": "user", "content": prompt},
        ])
        return MatchResult(**data)
